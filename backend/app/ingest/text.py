"""Steps 1-3: load file, extract text, chunk."""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


class ExtractionError(ValueError):
    pass


# --------------------------------------------------------------------------- extraction
def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ExtractionError(f"Unsupported file type {ext}. Supported: {', '.join(sorted(SUPPORTED))}")

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = [(p.extract_text() or "") for p in reader.pages]
        text = "\n\n".join(pages)
    elif ext == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:  # keep table content, row by row
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    else:
        text = path.read_text(encoding="utf-8", errors="replace")

    text = _clean(text)
    if len(text) < 20:
        raise ExtractionError(
            "No extractable text found (scanned/image-only PDF?). OCR is not supported in v1."
        )
    return text


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"-\n(?=[a-z])", "", text)        # de-hyphenate PDF line breaks
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# --------------------------------------------------------------------------- chunking
@dataclass
class Chunk:
    index: int
    text: str
    start: int  # char offset in the document text


_SEPARATORS = ["\n\n", "\n", ". ", " "]


def chunk_text(text: str, size: int, overlap: int) -> list[Chunk]:
    """Recursive splitter: prefer paragraph > line > sentence > word boundaries,
    then pack pieces into chunks of ~size chars with ~overlap chars carried over."""
    pieces = _split(text, size)
    chunks: list[Chunk] = []
    buf: list[str] = []
    buf_len = 0
    pos = 0
    buf_start = 0

    for piece in pieces:
        if buf and buf_len + len(piece) > size:
            chunk_str = "".join(buf).strip()
            if chunk_str:
                chunks.append(Chunk(len(chunks), chunk_str, buf_start))
            # carry overlap: keep trailing pieces up to `overlap` chars
            carry: list[str] = []
            carry_len = 0
            for p in reversed(buf):
                if carry_len + len(p) > overlap:
                    break
                carry.insert(0, p)
                carry_len += len(p)
            buf, buf_len = carry, carry_len
            buf_start = pos - carry_len
        if not buf:
            buf_start = pos
        buf.append(piece)
        buf_len += len(piece)
        pos += len(piece)

    tail = "".join(buf).strip()
    if tail:
        chunks.append(Chunk(len(chunks), tail, buf_start))
    return chunks


def _split(text: str, size: int, level: int = 0) -> list[str]:
    if len(text) <= size:
        return [text]
    if level >= len(_SEPARATORS):
        return [text[i : i + size] for i in range(0, len(text), size)]
    sep = _SEPARATORS[level]
    parts = text.split(sep)
    out: list[str] = []
    for i, part in enumerate(parts):
        piece = part + (sep if i < len(parts) - 1 else "")
        if len(piece) > size:
            out.extend(_split(piece, size, level + 1))
        elif piece:
            out.append(piece)
    return out
